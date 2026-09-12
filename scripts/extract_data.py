"""Read the supplied workbook into source-traceable simulator data. No workbook authoring."""
import json, re, sys, unicodedata, pathlib, operator
import openpyxl
sys.stdout.reconfigure(encoding='utf8')
ROOT=pathlib.Path(__file__).resolve().parent.parent / "source"
book=openpyxl.load_workbook(ROOT/'mabinogi-source.xlsx',data_only=False)
cached=openpyxl.load_workbook(ROOT/'mabinogi-source.xlsx',data_only=True)
BASE='https://docs.google.com/spreadsheets/d/1Ax4M2ycAv-XMDUo-DldQVt3cGiYRaz5cCTCONSoH46c/edit'
html=(ROOT/'source.html').read_text(encoding='utf8')
gids=dict(re.findall(r'items.push\(\{name: "([^"]+)",.*?gid: "(\d+)"',html))
ids=['blacksmith','fynn-bead-burnishing','tailoring','fynn-craft','enchant','magic-craft','hillwen-engineering','handicraft','stationery-craft']
names=['打鐵','芬恩寶珠閃耀','衣物製作','芬恩的手工藝','魔力賦予','魔法製造','稀原工學','手工藝','文具工藝']

class Formula:
 def __init__(self, text, resolve):
  self.tokens=re.findall(r'"(?:[^"]|"")*"|\$?[A-Za-z]+\$?\d+|[A-Za-z_]+|\d+(?:\.\d+)?|<>|<=|>=|[=<>+*/(),-]',text.lstrip('='))
  self.i=0;self.resolve=resolve
 def pop(self):
  t=self.tokens[self.i];self.i+=1;return t
 def peek(self):return self.tokens[self.i] if self.i<len(self.tokens) else None
 def expression(self,level=0):
  if level==3:return self.atom()
  groups=[['=','<>','<','>','<=','>='],['+','-'],['*','/']]
  a=self.expression(level+1)
  while self.peek() in groups[level]:
   op=self.pop();b=self.expression(level+1)
   fn={'=':operator.eq,'<>':operator.ne,'<':operator.lt,'>':operator.gt,'<=':operator.le,'>=':operator.ge,'+':operator.add,'-':operator.sub,'*':operator.mul,'/':operator.truediv}[op]
   a=fn(a,b)
  return a
 def atom(self):
  t=self.pop()
  if t=='-':return -self.atom()
  if t=='(':
   a=self.expression();assert self.pop()==')';return a
  if t.startswith('"'):return t[1:-1].replace('""','"')
  if re.match(r'^\d',t):return float(t)
  if self.peek()=='(':
   self.pop();args=[self.expression()]
   while self.peek()==',':self.pop();args.append(self.expression())
   assert self.pop()==')'
   if t.lower()=='if':return args[1] if args[0] else args[2]
   if t.lower()=='or':return any(args)
   raise ValueError('Unsupported function '+t)
  return self.resolve(t.replace('$',''))

def materials(s,row,recipe):
 result=[];unknown=[]
 def resolve(ref,visited=()):
  if ref in visited:raise ValueError('Circular reference')
  if ref==f'H{row}':return 1
  if ref==f'G{row}':return recipe
  if re.match(r'^H\d+$',ref):raise ValueError('References another task count '+ref)
  if re.findall(r'\d+',ref)[0]!=str(row):raise ValueError('References another source row '+ref)
  val=s[ref].value
  if val is None:return 0
  if isinstance(val,str) and val.startswith('='):
   return Formula(val,lambda x:resolve(x,visited+(ref,))).expression()
  return val
 for c in range(10, min(s.max_column,23)+1):
  name=s.cell(6,c).value
  if not name:continue
  cell=s.cell(row,c)
  if cell.value is None:continue
  try:
   qty=resolve(cell.coordinate)
   if qty in ('',None,0):continue
   if not isinstance(qty,(int,float)) or qty<0:raise ValueError('Non-positive material')
   result.append({'name':name,'perAction':qty,'sourceCell':cell.coordinate,'sourceFormula':cell.value,'basis':'source-formula-at-one-action'})
  except Exception as e:
   unknown.append({'name':name,'sourceCell':cell.coordinate,'sourceFormula':cell.value,'reason':str(e)})
 return result,unknown

skills=[];unresolved=[]
for s,sid,name in zip(book.worksheets[1:],ids,names):
 gid=gids.get(s.title)
 url=f'{BASE}?gid={gid}#gid={gid}' if gid else BASE
 headings=[(r,s.cell(r,1).value) for r in range(7,s.max_row+1) if isinstance(s.cell(r,1).value,str) and s.cell(r,2).value is None]
 skill={'id':sid,'name':name,'sourceSheet':s.title,'sourceUrl':url,'gid':gid,'sourceRange':s.calculate_dimension(),'apSummary':s['B1'].value,'notes':[],'defaults':{'countMultiplier':s['C2'].value,'valueMultiplier':s['C3'].value,'sourceCurrentRank':s['B4'].value,'sourceTarget':s['A3'].value},'ranks':[],'supplementaryRanks':[]}
 for cell in ['A5','G1','G2','H3']:
  v=s[cell].value
  if isinstance(v,str) and not v.startswith('='):
   skill['notes'].append({'text':v.strip(),'sourceCell':cell})
 for k,(start,label) in enumerate(headings):
  end=headings[k+1][0]-1 if k+1<len(headings) else s.max_row
  norm=unicodedata.normalize('NFKC',label).replace(' ','').strip()
  regular=norm.startswith('Rank')
  rank=norm.replace('Rank','') if regular else f'supplement-{start}'
  rdata={'id':f'{sid}-{rank}','rank':rank,'label':label,'sourceRange':f'A{start}:{openpyxl.utils.get_column_letter(min(s.max_column,23))}{end}','tasks':[]}
  taskrows=[r for r in range(start+1,end+1) if isinstance(s.cell(r,1).value,str) and isinstance(s.cell(r,2).value,(int,float)) and isinstance(s.cell(r,3).value,(int,float))]
  for ti,r in enumerate(taskrows):
   recipe=s.cell(r,7).value
   recipe=recipe.strip() if isinstance(recipe,str) else None
   mats,unknown=materials(s,r,recipe)
   hv=s.cell(r,8).value
   task={'id':f'{sid}-{rank}-{r}','description':s.cell(r,1).value.lstrip('☆').strip(),'baseValue':s.cell(r,2).value,'maxCount':s.cell(r,3).value,'recipe':recipe,'materials':mats,'materialStatus':'partial' if unknown else ('provided' if mats else 'not-provided'),'unresolvedMaterials':unknown,'sourceRow':r,'sourceRange':f'A{r}:{openpyxl.utils.get_column_letter(min(s.max_column,23))}{r}','recommended':isinstance(hv,str) and hv.startswith('=') or isinstance(hv,(int,float)) and hv>0,'sourcePlanFormula':hv if isinstance(hv,str) and hv.startswith('=') else None,'sourcePlanCount':cached[s.title].cell(r,8).value,'recipeVariants':[]}
   task['outcome']='failure' if any(t in task['description'] for t in ['失敗','不好']) else ('great-success' if any(t in task['description'] for t in ['大成功','非常成功','非常好']) else ('finish' if '完成' in task['description'] else ('attempt' if '嘗試' in task['description'] else 'success')))
   for dv in s.data_validations.dataValidation:
    if s.cell(r,7).coordinate in dv.sqref and dv.formula1 and dv.formula1.startswith('"'):
     for variant in dv.formula1[1:-1].split(','):
      vm,vu=materials(s,r,variant)
      task['recipeVariants'].append({'recipe':variant,'sourceRow':r,'materials':vm,'unresolvedMaterials':vu})
   stop=taskrows[ti+1] if ti+1<len(taskrows) else end+1
   for extra in range(r+1,stop):
    v=s.cell(extra,7).value
    if not v or s.cell(extra,5).value=='加總 :':continue
    vm,vu=materials(s,extra,v)
    task['recipeVariants'].append({'recipe':str(v).strip(),'sourceRow':extra,'materials':vm,'unresolvedMaterials':vu})
   if task['recipeVariants']:
    task['notes']=['同一修練項目的配方共用次數上限；列出的材料按選用配方計算。']
   if unknown:unresolved.append({'skill':sid,'rank':rank,'taskId':task['id'],'materials':unknown})
   rdata['tasks'].append(task)
  if rdata['tasks']:
   # Tight range: heading through the last row belonging to a task or its variants.
   finalrow=max([t['sourceRow'] for t in rdata['tasks']]+[v['sourceRow'] for t in rdata['tasks'] for v in t['recipeVariants']])
   rdata['sourceRange']=f'A{start}:{openpyxl.utils.get_column_letter(min(s.max_column,23))}{finalrow}'
   skill['ranks' if regular else 'supplementaryRanks'].append(rdata)
 skills.append(skill)

settings=book.worksheets[0]
count=[]
for row,ident in [(3,'talent'),(4,'training-potion'),(5,'royal-potion'),(6,'title'),(7,'constellation-balloon'),(8,'blue-prism'),(9,'pet')]:
 count.append({'id':ident,'name':settings.cell(row,1).value,'multiplier':settings.cell(row,3).value,'defaultEnabled':settings.cell(row,2).value,'sourceRange':f'A{row}:C{row}'})
output={'schemaVersion':1,'source':{'title':'瑪奇生活技能修練試算表','url':BASE+'?gid=785953580#gid=785953580','exportUrl':BASE.replace('/edit','/export?format=xlsx'),'retrievedAt':'2026-09-12','lastUpdated':'2025-02-04','lastUpdatedSourceCell':'說明&倍率計算!B26','authorCredit':'噴嚏精靈','authorCreditSourceCell':'說明&倍率計算!B22','scope':'9 skills from supplied workbook; not all Mabinogi life skills'},'multipliers':{'count':{'cap':8,'options':count,'eventDefault':2,'formula':'min(8, product(enabled bonuses) * event)','sourceRange':'說明&倍率計算!A2:C11'},'value':{'reforge':{'formula':'(10 + level) / 10','none':1,'heroBadge':1.5,'sourceRange':'說明&倍率計算!D3:F3'},'homestead':{'name':'浪漫農場／守護公會／訓練所','default':1,'sourceRange':'說明&倍率計算!D5:F5'},'pet':{'multiplier':2,'defaultEnabled':False,'sourceRange':'說明&倍率計算!D9:F9'},'eventDefault':1,'formula':'reforgeOrEchoOrBadge * homesteadOrGuildOrTraining * pet * event','sourceRange':'說明&倍率計算!D2:F11'}},'calculationNotes':['每項修練值 = baseValue × valueMultiplier × min(maxCount, 實際達成次數 × countMultiplier)。','次數倍率在來源公式上限為 8；修練值倍率不改變該項次數上限。','修練次數必須取整數。若距離目標仍有差額而該項已達上限，必須改用其他修練項目。','來源的材料總計常對含小計的範圍加總後除以 2，是排除重複加總；不可解讀為製作消耗減半。','材料按原表配方的指定達成次數估算；成功、失敗、大成功等結果不能保證，每次嘗試的實際耗材可能不同。','未提供的配方材料應顯示未提供，不能當作消耗為 0。','魔法製造及稀原工學的原表包含中間產物再利用及多配方路線；本資料保留每項直接配方，不把中間產物自動展開或抵扣。','原表附帶30%模式及AP修練提醒。此數據未提供逐rank AP費用，不可推算完整AP消耗。'],'skills':skills,'unresolvedMaterials':unresolved}
(ROOT/'skills-data.json').write_text(json.dumps(output,ensure_ascii=False,indent=2),encoding='utf8')
print('Wrote',ROOT/'skills-data.json')
print(json.dumps({'skillCount':len(skills),'rankCount':sum(len(s['ranks']) for s in skills),'taskCount':sum(len(r['tasks']) for s in skills for r in s['ranks']),'supplementaryCount':sum(len(s['supplementaryRanks']) for s in skills),'unresolvedCount':len(unresolved),'summary':[{'id':s['id'],'name':s['name'],'ranks':len(s['ranks']),'tasks':sum(len(r['tasks']) for r in s['ranks']),'gid':s['gid']} for s in skills]},ensure_ascii=False,indent=2))


